import sys
import io
import os
import shutil
import numpy as np
import cv2
from pathlib import Path
import fitz  # PyMuPDF
from pptx import Presentation
from fastapi.testclient import TestClient

# Force UTF-8 encoding
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

backend_dir = Path(__file__).resolve().parent / "backend"
if str(backend_dir) not in sys.path:
    sys.path.insert(0, str(backend_dir))

from backend.config import DATA_DIR, UPLOAD_DIR, MEDIA_DIR
from backend.database import (
    init_db, save_textbook, get_textbook, get_all_textbooks,
    get_chunks_for_textbook, record_assessment_attempt,
    upsert_misconception_profile, record_user_concept_progress,
    get_user_concept_progress
)
from backend.services.multimodal_service import process_multimodal_file
from backend.services.concept_graph_service import generate_course_concept_graph
from backend.services.visualization_service import (
    generate_visual_explanation,
    sanitize_mermaid_label,
    detect_best_diagram_type
)
from backend.services.learning_path_service import (
    generate_personalized_learning_path,
    explain_learning_path_node
)
from backend.main import app

client = TestClient(app)

print("=" * 80)
print("TESTING INNOVATION 5: AI-GENERATED VISUAL EXPLANATIONS & PERSONALIZED LEARNING PATHS")
print("=" * 80)

init_db()

test_dir = DATA_DIR / "test_innovation5"
test_dir.mkdir(parents=True, exist_ok=True)

# ----------------------------------------------------------------------
# SETUP COURSE 1: Distributed Systems & Consensus (PDF + PPTX + MP4)
# ----------------------------------------------------------------------
print("\n[Setup] Creating dynamic multimodal test course materials...")

course1_id = "course_dist_sys_501"

# 1. PDF Document
pdf_path_1 = test_dir / "Distributed_Systems_Core.pdf"
doc1 = fitz.open()
p1 = doc1.new_page()
p1.insert_textbox(
    fitz.Rect(50, 50, 550, 750),
    """Chapter 1: Distributed Consensus Protocols
1.1 Distributed Transactions
Distributed transactions span multiple database nodes and require atomic commit guarantees.
1.2 Two-Phase Commit Protocol (2PC)
The Two-Phase Commit Protocol is an atomic commitment protocol for distributed databases.
Before understanding Two-Phase Commit Protocol, one must master Distributed Transactions.
Phase 1: Prepare Phase. The coordinator node sends a PREPARE message to all cohort participants.
Phase 2: Commit Phase. If all cohorts reply YES, the coordinator sends a GLOBAL-COMMIT command. If any cohort replies NO or times out, the coordinator sends a GLOBAL-ABORT command.""",
    fontsize=11
)
doc1.save(str(pdf_path_1))
doc1.close()

# 2. PPTX Slides
pptx_path_1 = test_dir / "Distributed_Systems_Slides.pptx"
prs1 = Presentation()
s1 = prs1.slides.add_slide(prs1.slide_layouts[1])
s1.shapes.title.text = "Raft Consensus Algorithm"
tf1 = s1.shapes.placeholders[1].text_frame
tf1.text = "Raft consensus protocol manages replicated state machines across clusters."
p_a = tf1.add_paragraph()
p_a.text = "Raft Consensus Algorithm requires knowledge of Two-Phase Commit Protocol and leader election."
prs1.save(str(pptx_path_1))

# 3. MP4 Video
mp4_path_1 = test_dir / "Distributed_Systems_Video.mp4"
fps = 30
duration_sec = 2
total_frames = fps * duration_sec
fourcc = cv2.VideoWriter_fourcc(*'mp4v')
out = cv2.VideoWriter(str(mp4_path_1), fourcc, fps, (640, 480))
for f in range(total_frames):
    img = np.zeros((480, 640, 3), dtype=np.uint8)
    img[:, :, 0] = 120
    cv2.putText(img, "Lecture: Paxos Consensus and Quorums", (20, 240), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 255, 255), 2)
    out.write(img)
out.release()

# Process multimodal files and save textbook chunks
meta_pdf, chunks_pdf, _ = process_multimodal_file(pdf_path_1, "Distributed_Systems_Core.pdf")
meta_ppt, chunks_ppt, _ = process_multimodal_file(pptx_path_1, "Distributed_Systems_Slides.pptx")
meta_mp4, chunks_mp4, _ = process_multimodal_file(mp4_path_1, "Distributed_Systems_Video.mp4")

all_chunks_1 = []
for c in chunks_pdf + chunks_ppt + chunks_mp4:
    c["textbook_id"] = course1_id
    all_chunks_1.append(c)

meta_course_1 = {
    "id": course1_id,
    "title": "Distributed Systems & Consensus Architecture",
    "filename": "Distributed_Systems_Core.pdf",
    "source_type": "multimodal",
    "total_pages": 1,
    "total_slides": 1,
    "duration_seconds": 2.0,
    "duration_str": "00:00:02",
    "file_size_bytes": 10240,
    "chapters": meta_pdf.get("chapters", []) + meta_ppt.get("chapters", [])
}

save_textbook(meta_course_1, all_chunks_1)
graph_c1 = generate_course_concept_graph(course1_id, force_regenerate=True)

print(f"Course 1 loaded: {len(all_chunks_1)} chunks, {len(graph_c1['nodes'])} concepts discovered.")


# ----------------------------------------------------------------------
# TEST 1: DYNAMIC FLOWCHART GENERATION FROM PROCESS IN SOURCE
# ----------------------------------------------------------------------
print("\n[Test 1] Dynamic Flowchart Generation for Two-Phase Commit Protocol...")

res_flow = generate_visual_explanation(
    chunks=all_chunks_1,
    topic="Two-Phase Commit Protocol",
    diagram_type="flowchart"
)

assert res_flow["detected_type"] == "flowchart", f"Expected flowchart, got {res_flow['detected_type']}"
assert "graph TD" in res_flow["mermaid_code"], "Mermaid code missing 'graph TD'"
assert "Decision" in res_flow["mermaid_code"] or "Check" in res_flow["mermaid_code"] or "DEC" in res_flow["mermaid_code"], "Flowchart missing decision branch"
assert res_flow["is_sufficient_evidence"] is True, "Evidence should be sufficient"
assert len(res_flow["steps_or_nodes"]) > 0, "Steps list should not be empty"

print(f"✓ Flowchart generated successfully:")
print(f"  Topic: {res_flow['topic']}")
print(f"  Type: {res_flow['detected_type']}")
print(f"  Steps Count: {len(res_flow['steps_or_nodes'])}")
print(f"  Mermaid Preview:\n{res_flow['mermaid_code'][:150]}...")


# ----------------------------------------------------------------------
# TEST 2: MULTI-FORMAT DIAGRAM GENERATION (Sequence, Block, Concept Map, Comparison, Timeline)
# ----------------------------------------------------------------------
print("\n[Test 2] Testing Multi-Format Diagram Generation...")

# 2a. Sequence Diagram (Protocol interaction)
res_seq = generate_visual_explanation(
    chunks=all_chunks_1,
    topic="Two-Phase Commit Protocol",
    diagram_type="sequence"
)
assert res_seq["detected_type"] == "sequence", f"Expected sequence, got {res_seq['detected_type']}"
assert "sequenceDiagram" in res_seq["mermaid_code"], "Mermaid code missing 'sequenceDiagram'"
assert "participant" in res_seq["mermaid_code"], "Sequence diagram missing participant declarations"
print("✓ Sequence diagram generated with message flow.")

# 2b. Block Diagram (Architecture)
res_block = generate_visual_explanation(
    chunks=all_chunks_1,
    topic="Distributed Consensus Protocols",
    diagram_type="block_diagram"
)
assert res_block["detected_type"] == "block_diagram"
assert "subgraph" in res_block["mermaid_code"]
print("✓ Block diagram generated with modular architecture subgraphs.")

# 2c. Concept Map (Semantic Network)
res_cmap = generate_visual_explanation(
    chunks=all_chunks_1,
    topic="Distributed Consensus Protocols",
    diagram_type="concept_map"
)
assert res_cmap["detected_type"] == "concept_map"
assert "-->|" in res_cmap["mermaid_code"] or "-->" in res_cmap["mermaid_code"]
print("✓ Concept map generated with relational edges.")

# 2d. Comparison Diagram
res_comp = generate_visual_explanation(
    chunks=all_chunks_1,
    topic="Two-Phase Commit vs Raft Consensus",
    diagram_type="comparison"
)
assert res_comp["detected_type"] == "comparison"
assert "subgraph" in res_comp["mermaid_code"]
print("✓ Comparison diagram generated with contrasting subgraphs.")

# 2e. Timeline Diagram
res_time = generate_visual_explanation(
    chunks=all_chunks_1,
    topic="Two-Phase Commit Protocol",
    diagram_type="timeline"
)
assert res_time["detected_type"] == "timeline"
print("✓ Timeline diagram generated.")


# ----------------------------------------------------------------------
# TEST 3: EVIDENCE GROUNDING & ACTUAL SOURCE CITATIONS
# ----------------------------------------------------------------------
print("\n[Test 3] Verifying Evidence Grounding and Provenance Citations...")

for step in res_flow["steps_or_nodes"]:
    cit = step.get("source_citation")
    assert cit is not None, f"Step '{step['label']}' missing source citation!"
    assert cit.get("filename") is not None, "Citation missing filename"
    assert (cit.get("page") is not None or cit.get("slide") is not None or cit.get("timestamp") is not None), "Citation missing page/slide/timestamp"
    assert step.get("is_supported") is True

print(f"✓ All {len(res_flow['steps_or_nodes'])} nodes have verified grounded provenance citations.")
print(f"  Sample Citation: File='{res_flow['steps_or_nodes'][0]['source_citation']['filename']}', Page={res_flow['steps_or_nodes'][0]['source_citation'].get('page')}")


# ----------------------------------------------------------------------
# TEST 4: INSUFFICIENT EVIDENCE HANDLING (Honest Limitation Notice)
# ----------------------------------------------------------------------
print("\n[Test 4] Testing Insufficient Evidence Detection...")

res_insuf = generate_visual_explanation(
    chunks=all_chunks_1,
    topic="Quantum Entanglement and Superconductivity",
    diagram_type="auto"
)

assert res_insuf["is_sufficient_evidence"] is False, "Expected insufficient evidence flag to be False!"
assert "Insufficient" in res_insuf["description"] or "insufficient" in res_insuf["limitation_note"].lower(), "Limitation explanation missing!"
assert len(res_insuf["steps_or_nodes"]) == 0, "Should not hallucinate fake steps when evidence is insufficient"
assert "⚠️" in res_insuf["mermaid_code"] or "Insufficient" in res_insuf["mermaid_code"], "Mermaid should render honest notice banner"

print(f"✓ Insufficient evidence handled honestly without fabricating diagrams:")
print(f"  is_sufficient_evidence: {res_insuf['is_sufficient_evidence']}")
print(f"  Limitation Note: {res_insuf['limitation_note']}")


# ----------------------------------------------------------------------
# TEST 5: LONG LABELS & SPECIAL CHARACTERS SANITIZATION
# ----------------------------------------------------------------------
print("\n[Test 5] Testing Long Labels & Safe Mermaid Sanitization...")

messy_label = 'Node With [Special] {Brackets} "Quotes" <script>alert("xss")</script> and very very long descriptive educational text that exceeds sixty characters'
clean_lbl = sanitize_mermaid_label(messy_label, max_length=50)

assert '"' not in clean_lbl, "Unescaped double quotes should be removed"
assert '[' not in clean_lbl and ']' not in clean_lbl, "Square brackets should be sanitized"
assert '<script>' not in clean_lbl, "HTML script tags should be stripped"
assert len(clean_lbl) <= 50, f"Label length {len(clean_lbl)} exceeds max limit 50"

print(f"✓ Sanitization result: '{clean_lbl}'")


# ----------------------------------------------------------------------
# TEST 6: SELECTED NODE EXPLANATION THROUGH GROUNDED TUTOR
# ----------------------------------------------------------------------
print("\n[Test 6] Testing Selected Node Simpler Explanation via Grounded Tutor...")

node_exp = explain_learning_path_node(
    course_id=course1_id,
    topic="Two-Phase Commit Protocol",
    node_label="Prepare Phase",
    language="en"
)

assert node_exp["node_label"] == "Prepare Phase"
assert len(node_exp["simple_explanation"]) > 20, "Simple explanation should be informative"
assert node_exp["source_citation"] is not None, "Source citation should be attached"

print(f"✓ Grounded Node Explanation:")
print(f"  Explanation: {node_exp['simple_explanation'][:120]}...")
if node_exp.get("analogy"):
    print(f"  Analogy: {node_exp['analogy']}")


# ----------------------------------------------------------------------
# TEST 7: PREREQUISITE-BASED TOPOLOGICAL LEARNING PATH GENERATION
# ----------------------------------------------------------------------
print("\n[Test 7] Testing Prerequisite-Based Topological Learning Path Generation...")

# Generate path for "Raft Consensus Algorithm" (Target Concept)
path_res = generate_personalized_learning_path(
    course_id=course1_id,
    target_concept="Raft Consensus Algorithm",
    user_id="user_untested_101"
)

assert path_res["target_concept"] == "Raft Consensus Algorithm"
assert path_res["total_steps"] >= 2, "Path should contain prerequisite steps"
assert path_res["is_personalized_with_assessment"] is False, "Untested user should have cold-start flag"
assert "Standard Prerequisite Pathway" in path_res["recommendation_rationale"], "Cold start notice should be present"

step_names = [s["concept_name"] for s in path_res["steps"]]
print(f"✓ Topological Learning Path generated:")
print(f"  Steps Sequence: {' -> '.join(step_names)}")
print(f"  Total Steps: {path_res['total_steps']}")
print(f"  Rationale: {path_res['recommendation_rationale']}")


# ----------------------------------------------------------------------
# TEST 8: ADAPTIVE PERSONALIZATION VIA ASSESSMENT RESULTS & MISCONCEPTIONS
# ----------------------------------------------------------------------
print("\n[Test 8] Testing Assessment Performance & Misconception Integration into Learning Path...")

student_user_id = "student_adaptive_999"

# 1. Student masters "Distributed Transactions" with 100% score (2 correct attempts)
record_assessment_attempt({
    "user_id": student_user_id,
    "course_id": course1_id,
    "question_id": "q_test_1",
    "question_text": "What is a distributed transaction?",
    "topic": "Distributed Transactions",
    "concept_name": "Distributed Transactions",
    "is_correct": 1,
    "score": 1.0,
    "evaluation_status": "correct",
    "feedback": "Correct definition.",
    "difficulty": "easy"
})
record_assessment_attempt({
    "user_id": student_user_id,
    "course_id": course1_id,
    "question_id": "q_test_2",
    "question_text": "How do distributed transactions ensure atomicity?",
    "topic": "Distributed Transactions",
    "concept_name": "Distributed Transactions",
    "is_correct": 1,
    "score": 1.0,
    "evaluation_status": "correct",
    "feedback": "Correct.",
    "difficulty": "medium"
})

# 2. Student fails "Two-Phase Commit Protocol" and is diagnosed with an active misconception
record_assessment_attempt({
    "user_id": student_user_id,
    "course_id": course1_id,
    "question_id": "q_test_3",
    "question_text": "Can a cohort commit independently in Phase 1 of 2PC?",
    "topic": "Two-Phase Commit Protocol",
    "concept_name": "Two-Phase Commit Protocol",
    "is_correct": 0,
    "score": 0.0,
    "evaluation_status": "incorrect",
    "feedback": "Cohorts must wait for Global-Commit in Phase 2.",
    "difficulty": "medium"
})

upsert_misconception_profile(
    user_id=student_user_id,
    course_id=course1_id,
    misconception={
        "concept_name": "Two-Phase Commit Protocol",
        "topic": "Two-Phase Commit Protocol",
        "misconception_title": "Premature Independent Cohort Commit",
        "hypothesis_text": "Student believes cohorts commit locally during the Prepare phase without waiting for the Coordinator Global Commit message.",
        "expected_understanding": "Cohorts only promise readiness in Phase 1 and must wait for Phase 2 Global-Commit.",
        "confidence": "high"
    }
)

# 3. Generate Personalized Learning Path for this student
personalized_path = generate_personalized_learning_path(
    course_id=course1_id,
    target_concept="Raft Consensus Algorithm",
    user_id=student_user_id
)

assert personalized_path["is_personalized_with_assessment"] is True, "Should be marked personalized with assessment"
assert len(personalized_path["active_misconceptions"]) > 0, "Should contain active diagnosed misconceptions"

# Verify concept statuses
dist_tx_step = next((s for s in personalized_path["steps"] if "Distributed Transactions" in s["concept_name"]), None)
two_pc_step = next((s for s in personalized_path["steps"] if "Two-Phase Commit" in s["concept_name"]), None)

assert dist_tx_step is not None and dist_tx_step["mastery_status"] == "mastered", f"Expected Dist Tx to be mastered, got {dist_tx_step['mastery_status'] if dist_tx_step else None}"
assert two_pc_step is not None and two_pc_step["mastery_status"] == "needs_revision", f"Expected 2PC to be needs_revision, got {two_pc_step['mastery_status'] if two_pc_step else None}"
assert two_pc_step["has_active_misconception"] is True, "2PC should flag active misconception"

# Verify Next Recommendation points to reviewing the struggling prerequisite
assert "Two-Phase Commit" in personalized_path["next_recommended_concept"], f"Expected next concept to be 2PC revision, got '{personalized_path['next_recommended_concept']}'"
assert personalized_path["recommended_action_type"] in ["review_prerequisite", "resolve_misconception"], f"Expected review/remediation action, got {personalized_path['recommended_action_type']}"

print("✓ Adaptive personalization verified:")
print(f"  Is Personalized: {personalized_path['is_personalized_with_assessment']}")
print(f"  Step 1 (Distributed Transactions): Status={dist_tx_step['mastery_status']}")
print(f"  Step 2 (Two-Phase Commit): Status={two_pc_step['mastery_status']}, Misconception='{two_pc_step['active_misconception_title']}'")
print(f"  Next Recommended Action: {personalized_path['recommended_action_type']} on '{personalized_path['next_recommended_concept']}'")
print(f"  Rationale: {personalized_path['recommendation_rationale']}")


# ----------------------------------------------------------------------
# TEST 9: STUDENT STUDY PROGRESS TRACKING
# ----------------------------------------------------------------------
print("\n[Test 9] Testing Student Self-Directed Study Progress Checkpoints...")

record_user_concept_progress(
    user_id=student_user_id,
    course_id=course1_id,
    concept_id=two_pc_step["concept_id"],
    concept_name="Two-Phase Commit Protocol",
    status="studied",
    notes="Reviewed visual diagram and coordinator commit steps."
)

prog_items = get_user_concept_progress(student_user_id, course1_id)
assert len(prog_items) > 0
assert prog_items[0]["status"] == "studied"
assert "Reviewed visual diagram" in prog_items[0]["notes"]

print(f"✓ Manual concept progress saved and retrieved ({len(prog_items)} records).")


# ----------------------------------------------------------------------
# TEST 10: FASTAPI HTTP ENDPOINTS TEST (Innovation 5)
# ----------------------------------------------------------------------
print("\n[Test 10] Testing Innovation 5 HTTP Endpoints via TestClient...")

# 1. /api/visualize
resp_vis = client.post("/api/visualize", json={
    "textbook_id": course1_id,
    "topic": "Two-Phase Commit Protocol",
    "diagram_type": "sequence"
})
assert resp_vis.status_code == 200
vis_json = resp_vis.json()
assert vis_json["detected_type"] == "sequence"
assert "mermaid_code" in vis_json

# 2. /api/learning-path/personalized (POST)
resp_lp_post = client.post("/api/learning-path/personalized", json={
    "course_id": course1_id,
    "target_concept": "Raft Consensus Algorithm",
    "user_id": student_user_id
})
assert resp_lp_post.status_code == 200
lp_json = resp_lp_post.json()
assert lp_json["target_concept"] == "Raft Consensus Algorithm"
assert len(lp_json["steps"]) >= 2

# 3. /api/learning-path/personalized/{course_id} (GET)
resp_lp_get = client.get(f"/api/learning-path/personalized/{course1_id}?user_id={student_user_id}")
assert resp_lp_get.status_code == 200

# 4. /api/learning-path/progress (POST & GET)
resp_prog_post = client.post("/api/learning-path/progress", json={
    "course_id": course1_id,
    "concept_id": "c_test_01",
    "concept_name": "Paxos",
    "status": "studied",
    "user_id": student_user_id
})
assert resp_prog_post.status_code == 200

resp_prog_get = client.get(f"/api/learning-path/progress/{course1_id}?user_id={student_user_id}")
assert resp_prog_get.status_code == 200
assert len(resp_prog_get.json()["progress_items"]) >= 1

# 5. /api/learning-path/node-explain (POST)
resp_node = client.post("/api/learning-path/node-explain", json={
    "course_id": course1_id,
    "topic": "Two-Phase Commit Protocol",
    "node_label": "Commit Phase"
})
assert resp_node.status_code == 200
assert "simple_explanation" in resp_node.json()

print("✓ All 5 Innovation HTTP endpoints tested and passed.")


# ----------------------------------------------------------------------
# TEST 11: REGRESSION TESTING INNOVATIONS 1 - 4
# ----------------------------------------------------------------------
print("\n[Test 11] Regression Testing Innovations 1, 2, 3, and 4...")

# Regression Innovation 1: Multimodal Textbook list
tb_list = client.get("/api/textbooks").json()
assert len(tb_list) >= 1
print("✓ Innovation 1 (Multimodal Ingestion): OK")

# Regression Innovation 2: Concept Graph
cg_res = client.get(f"/api/concept-graph/{course1_id}").json()
assert len(cg_res["nodes"]) >= 1
assert len(cg_res["edges"]) >= 0
print("✓ Innovation 2 (Concept Graph & Prerequisites): OK")

# Regression Innovation 3: Grounded Tutor RAG
chat_res = client.post("/api/chat", json={
    "textbook_id": course1_id,
    "query": "What is Phase 1 of Two-Phase Commit Protocol?"
}).json()
assert "citations" in chat_res
assert len(chat_res["citations"]) > 0
print("✓ Innovation 3 (Anti-Hallucination Grounded Tutor): OK")

# Regression Innovation 4: Adaptive Assessment
assess_res = client.post("/api/assessment/generate", json={
    "course_id": course1_id,
    "topic": "All",
    "difficulty": "medium",
    "num_questions": 3
}).json()
assert len(assess_res["questions"]) > 0
print("✓ Innovation 4 (Misconception-Aware Adaptive Assessment): OK")

print("\n" + "=" * 80)
print("ALL INNOVATION 5 TESTS PASSED SUCCESSFULLY (11/11 TEST SUITES)")
print("=" * 80)
