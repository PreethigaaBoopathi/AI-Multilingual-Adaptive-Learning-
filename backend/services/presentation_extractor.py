import os
from pathlib import Path
from typing import Dict, Any, List, Tuple, Optional
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from config import MEDIA_DIR

def extract_from_presentation(pptx_path: Path, filename: str, custom_id: Optional[str] = None) -> Tuple[Dict[str, Any], List[Dict[str, Any]], List[Dict[str, Any]]]:
    """
    Extracts slide text, slide numbers, images, tables, speaker notes, and structure
    dynamically from any PPTX lecture presentation.
    Preserves exact slide-level provenance for every content unit.
    """
    if not pptx_path.exists():
        raise FileNotFoundError(f"Presentation file not found: {pptx_path}")

    try:
        prs = Presentation(str(pptx_path))
    except Exception as e:
        raise ValueError(f"Unable to parse presentation file: {str(e)}. Please ensure the presentation is a valid PowerPoint file.")

    total_slides = len(prs.slides)
    if total_slides == 0:
        raise ValueError("The presentation file contains 0 slides.")

    file_size_bytes = pptx_path.stat().st_size
    file_id = custom_id or f"ppt_{abs(hash(filename + str(file_size_bytes))) % 1000000}"
    media_dir = MEDIA_DIR / file_id
    media_dir.mkdir(parents=True, exist_ok=True)

    clean_title = filename.replace(".pptx", "").replace(".ppt", "").replace(".PPTX", "").replace(".PPT", "").replace("_", " ").replace("-", " ").title()

    slides_data = []
    media_items = []
    extracted_tables_count = 0

    # 1. Slide-by-slide extraction
    for slide_idx, slide in enumerate(prs.slides):
        slide_num = slide_idx + 1
        slide_title = ""
        slide_texts = []
        slide_tables = []
        speaker_notes = ""

        # Check for slide title shape
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            slide_title = slide.shapes.title.text_frame.text.strip()

        # Iterate over all shapes on the slide
        img_counter = 1
        for shape in slide.shapes:
            # Extract text from shapes
            if shape.has_text_frame and shape != slide.shapes.title:
                text = shape.text_frame.text.strip()
                if text:
                    slide_texts.append(text)

            # Extract tables
            if shape.has_table:
                table = shape.table
                table_lines = []
                num_cols = len(table.columns)
                # Extract header / rows
                for r_idx, row in enumerate(table.rows):
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    table_lines.append(" | ".join(row_cells))
                    if r_idx == 0:
                        table_lines.append(" | ".join(["---"] * num_cols))
                if table_lines:
                    table_md = "\n".join(table_lines)
                    slide_tables.append(table_md)
                    extracted_tables_count += 1

            # Extract embedded pictures / figures
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = shape.image
                    img_bytes = image.blob
                    img_ext = image.ext or "png"
                    img_filename = f"slide_{slide_num}_img_{img_counter}.{img_ext}"
                    img_path = media_dir / img_filename

                    with open(img_path, "wb") as f_img:
                        f_img.write(img_bytes)

                    media_url = f"/media/{file_id}/{img_filename}"
                    media_item = {
                        "id": f"med_{file_id}_s{slide_num}_{img_counter}",
                        "file_id": file_id,
                        "filename": filename,
                        "source_type": "presentation",
                        "content_type": "image",
                        "page": slide_num,
                        "slide": slide_num,
                        "timestamp": None,
                        "media_path": media_url,
                        "media_filename": img_filename,
                        "format": img_ext,
                        "description": f"Visual diagram/figure embedded in Slide {slide_num} of {filename}"
                    }
                    media_items.append(media_item)
                    img_counter += 1
                except Exception:
                    pass

        # Extract Speaker Notes if present
        try:
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide.notes_text_frame:
                    notes_txt = notes_slide.notes_text_frame.text.strip()
                    if notes_txt:
                        speaker_notes = notes_txt
        except Exception:
            pass

        # Fallback slide title if not explicitly set
        if not slide_title:
            if slide_texts:
                first_line = slide_texts[0].split("\n")[0].strip()
                slide_title = first_line[:60] if len(first_line) > 3 else f"Slide {slide_num}"
            else:
                slide_title = f"Slide {slide_num}"

        slides_data.append({
            "slide_num": slide_num,
            "title": slide_title,
            "texts": slide_texts,
            "tables": slide_tables,
            "speaker_notes": speaker_notes
        })

    # Validate that presentation has content
    has_text = any(s["texts"] or s["title"] != f"Slide {s['slide_num']}" or s["tables"] or s["speaker_notes"] for s in slides_data)
    if not has_text and not media_items:
        raise ValueError("Unable to extract content from this presentation. The slides contain no readable text or media.")

    # 2. Build chapter/topic structure based on slides
    # If first slide looks like a title slide, use its text as presentation title
    if slides_data and slides_data[0]["title"] and slides_data[0]["title"] != "Slide 1":
        clean_title = slides_data[0]["title"]

    chapters = []
    slides_per_ch = max(1, (total_slides + 3) // 4)
    
    for i in range(0, total_slides, slides_per_ch):
        ch_n = (i // slides_per_ch) + 1
        s_st = i + 1
        s_en = min(i + slides_per_ch, total_slides)
        first_slide_in_group = slides_data[i]
        ch_title = f"Topic {ch_n}: {first_slide_in_group['title']}"

        sections = []
        for s_idx in range(s_st - 1, s_en):
            s_item = slides_data[s_idx]
            sections.append({
                "id": f"sec_{ch_n}_{s_item['slide_num']}",
                "title": f"Slide {s_item['slide_num']}: {s_item['title']}",
                "page_start": s_item["slide_num"],
                "page_end": s_item["slide_num"],
                "content": ""
            })

        chapters.append({
            "id": f"ch_{ch_n}",
            "number": ch_n,
            "title": ch_title,
            "page_start": s_st,
            "page_end": s_en,
            "sections": sections
        })

    # 3. Create unified knowledge chunks with exact slide provenance
    chunks = []
    chunk_counter = 1

    for s_info in slides_data:
        s_num = s_info["slide_num"]
        s_title = s_info["title"]

        # Match to chapter
        matched_ch = chapters[0]
        for ch in chapters:
            if ch["page_start"] <= s_num <= ch["page_end"]:
                matched_ch = ch
                break

        # Slide Main Text Chunks
        full_slide_content = f"Slide {s_num}: {s_title}\n" + "\n".join(s_info["texts"])
        if s_info["texts"]:
            for t_block in s_info["texts"]:
                chunks.append({
                    "id": f"chk_{file_id}_{chunk_counter}",
                    "textbook_id": file_id,
                    "material_id": file_id,
                    "source_type": "presentation",
                    "filename": filename,
                    "chapter_id": matched_ch["id"],
                    "chapter_title": matched_ch["title"],
                    "section_id": f"sec_{matched_ch['number']}_{s_num}",
                    "section_title": f"Slide {s_num}: {s_title}",
                    "page": s_num,
                    "slide_number": s_num,
                    "start_time": None,
                    "end_time": None,
                    "timestamp_str": None,
                    "content_type": "text",
                    "content": f"[Slide {s_num} - {s_title}]: {t_block}",
                    "media_path": None,
                    "metadata": {
                        "source": "slide_text",
                        "slide_number": s_num,
                        "slide_title": s_title,
                        "filename": filename
                    }
                })
                chunk_counter += 1
        else:
            # Even if only title exists
            chunks.append({
                "id": f"chk_{file_id}_{chunk_counter}",
                "textbook_id": file_id,
                "material_id": file_id,
                "source_type": "presentation",
                "filename": filename,
                "chapter_id": matched_ch["id"],
                "chapter_title": matched_ch["title"],
                "section_id": f"sec_{matched_ch['number']}_{s_num}",
                "section_title": f"Slide {s_num}: {s_title}",
                "page": s_num,
                "slide_number": s_num,
                "start_time": None,
                "end_time": None,
                "timestamp_str": None,
                "content_type": "text",
                "content": f"[Slide {s_num}]: {s_title}",
                "media_path": None,
                "metadata": {
                    "source": "slide_title",
                    "slide_number": s_num,
                    "slide_title": s_title,
                    "filename": filename
                }
            })
            chunk_counter += 1

        # Slide Table Chunks
        for tbl_str in s_info["tables"]:
            chunks.append({
                "id": f"chk_{file_id}_tbl_{chunk_counter}",
                "textbook_id": file_id,
                "material_id": file_id,
                "source_type": "presentation",
                "filename": filename,
                "chapter_id": matched_ch["id"],
                "chapter_title": matched_ch["title"],
                "section_id": f"sec_{matched_ch['number']}_{s_num}",
                "section_title": f"Slide {s_num} Table: {s_title}",
                "page": s_num,
                "slide_number": s_num,
                "start_time": None,
                "end_time": None,
                "timestamp_str": None,
                "content_type": "table",
                "content": f"[Slide {s_num} Table - {s_title}]:\n{tbl_str}",
                "media_path": None,
                "metadata": {
                    "source": "slide_table",
                    "slide_number": s_num,
                    "filename": filename
                }
            })
            chunk_counter += 1

        # Speaker Notes Chunks
        if s_info["speaker_notes"]:
            chunks.append({
                "id": f"chk_{file_id}_note_{chunk_counter}",
                "textbook_id": file_id,
                "material_id": file_id,
                "source_type": "presentation",
                "filename": filename,
                "chapter_id": matched_ch["id"],
                "chapter_title": matched_ch["title"],
                "section_id": f"sec_{matched_ch['number']}_{s_num}",
                "section_title": f"Slide {s_num} Speaker Notes",
                "page": s_num,
                "slide_number": s_num,
                "start_time": None,
                "end_time": None,
                "timestamp_str": None,
                "content_type": "speaker_notes",
                "content": f"[Slide {s_num} Speaker Notes]: {s_info['speaker_notes']}",
                "media_path": None,
                "metadata": {
                    "source": "slide_speaker_notes",
                    "slide_number": s_num,
                    "filename": filename
                }
            })
            chunk_counter += 1

    # Add embedded images as content units
    for img_item in media_items:
        i_slide = img_item["slide"]
        matched_ch = chapters[0]
        for ch in chapters:
            if ch["page_start"] <= i_slide <= ch["page_end"]:
                matched_ch = ch
                break

        chunks.append({
            "id": f"chk_{file_id}_img_{chunk_counter}",
            "textbook_id": file_id,
            "material_id": file_id,
            "source_type": "presentation",
            "filename": filename,
            "chapter_id": matched_ch["id"],
            "chapter_title": matched_ch["title"],
            "section_id": f"sec_{matched_ch['number']}_{i_slide}",
            "section_title": f"Slide {i_slide} Embedded Figure",
            "page": i_slide,
            "slide_number": i_slide,
            "start_time": None,
            "end_time": None,
            "timestamp_str": None,
            "content_type": "image",
            "content": f"[Embedded Figure / Visual Diagram on Slide {i_slide} of {filename}]",
            "media_path": img_item["media_path"],
            "metadata": {
                "source": "slide_embedded_image",
                "slide_number": i_slide,
                "filename": filename,
                "image_url": img_item["media_path"]
            }
        })
        chunk_counter += 1

    metadata = {
        "id": file_id,
        "title": clean_title,
        "filename": filename,
        "source_type": "presentation",
        "total_pages": total_slides,
        "total_slides": total_slides,
        "duration_seconds": None,
        "file_size_bytes": file_size_bytes,
        "extracted_images_count": len(media_items),
        "extracted_transcripts_count": 0,
        "extracted_tables_count": extracted_tables_count,
        "media_items": media_items,
        "chapters": chapters,
        "created_at": "2026-10-07T12:00:00"
    }

    return metadata, chunks, media_items
